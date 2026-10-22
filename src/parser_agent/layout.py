"""Deterministic detection of a sheet header: where the header ends and the data
begins. Without an LLM — from the structure of the sheet alone.

Two table orientations are distinguished.

**Orientation 1 — staggered.** The columns start at different rows: the left
label column starts below the others. Everything above its first row is the
header, however many rows it takes. A column title is glued together from all
of its cells in that zone. The label column itself has no header cells, so its
name is taken from the sheet name.

    A       B      C          <- B is entirely empty
    .       .      C1  ┐
    .       .      C2  │ header of column C: C1..C4
    .       .      C3  │
    .       .      C4  ┘
    A5      .      C5  <- the data starts here

An upper level of the header may be written without a merge range: it sits in
one column, and the columns to its right carry the levels below it. Such a value
spans to the right over every following column, no matter how many cells that
column's own header has, and stops at the first column that has a value of its
own in that row.

    C1 ────────────────▶       H1 ────▶     <- C1 spans C..G, H1 starts a new group
    C2      D2    E2 F2 G2     H2
    C3            E3 F3 G3     H3           <- D holds one cell, it is spanned too

Only captions are read from the header zone: a word, a date or a year. A figure
there is not a title — reports keep check totals above the table, and a header
formula pointing at an empty cell quietly returns 0. Neither may end up glued
into a column name.

**Orientation 2 — typed.** All columns start at the same row. Then the header
boundary is set by the data type change in the label column: the leading run of
cells of one type is the header, the first cell of another type is the start of
the data. Merged cells of the same type stay part of the header. In this case
the name of the label column is known from its own header cells, and the sheet
name is not needed.

    A1: string  ┐
    A2: string  │ header (3 rows)
    A3: string  ┘
    A4: integer <- data
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .compressor.aggregator import detect_type
from .models import sanitize_name_part

# The header level separator inside a chain. A double underscore, so that a
# level boundary differs from a space inside one level (a space becomes "_").
HEADER_LEVEL_SEPARATOR = "__"

# The same safety limits as in the compressor: Excel keeps formatting far
# beyond the real data.
MAX_ROWS = 20_000
MAX_COLS = 256

# How many columns the header chains cover — the list the model builds the
# schema from. It must hold every column of a real report: a column cut off
# here never reaches the database.
MAX_HEADER_COLS = 200

# Types considered one class when looking for a type change: for the header
# boundary what matters is «text vs non-text», not Integer vs Float.
_TEXT_LIKE = {"Text", "Email"}

# The value types a header cell may hold. Everything else is a figure, and a
# figure is never a column title.
_CAPTION_TYPES = frozenset({"Text", "Email", "Year", "Date"})


@dataclass
class HeaderLayout:
    """The result of parsing the sheet header."""

    orientation: str                 # "staggered" | "typed" | "unknown"
    header_rows: list[int] = field(default_factory=list)
    data_start_row: int | None = None
    data_end_row: int | None = None
    label_col: int | None = None     # the label column (leftmost column with data)
    label_has_header: bool = False   # the label column has header cells of its own
    label_name_from_sheet: bool = False  # take the label column name from the sheet name
    non_empty_cols: list[int] = field(default_factory=list)  # columns holding anything at all

    @property
    def detected(self) -> bool:
        return self.orientation != "unknown" and self.data_start_row is not None


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _merge_lookup(ws: Worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    """(row, column) -> the coordinates of the merge range anchor."""
    lookup: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        height = rng.max_row - rng.min_row + 1
        width = rng.max_col - rng.min_col + 1
        if height * width > 100_000:
            continue
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                lookup[(row, col)] = (rng.min_row, rng.min_col)
    return lookup


def cell_value(
    ws: Worksheet, row: int, col: int, merged: dict[tuple[int, int], tuple[int, int]]
) -> object:
    """The cell value, with the merge range expanded."""
    value = ws.cell(row=row, column=col).value
    if _is_blank(value) and (row, col) in merged:
        anchor_row, anchor_col = merged[(row, col)]
        value = ws.cell(row=anchor_row, column=anchor_col).value
    return value


def column_extents(
    ws: Worksheet, merged: dict[tuple[int, int], tuple[int, int]] | None = None
) -> dict[int, tuple[int, int]]:
    """{column: (first non-empty row, last non-empty row)}.

    Completely empty columns do not make it into the result: in reports they are
    used as block separators and must not affect the header parsing.
    """
    merged = _merge_lookup(ws) if merged is None else merged
    n_rows = min(ws.max_row or 1, MAX_ROWS)
    n_cols = min(ws.max_column or 1, MAX_COLS)

    extents: dict[int, tuple[int, int]] = {}
    for col in range(1, n_cols + 1):
        first = last = None
        for row in range(1, n_rows + 1):
            if not _is_blank(cell_value(ws, row, col, merged)):
                first = row if first is None else first
                last = row
        if first is not None:
            extents[col] = (first, last)
    return extents


def _type_class(value: object) -> str:
    """A coarse type class: for the header boundary text vs non-text is what matters."""
    kind = detect_type(value)
    return "text" if kind in _TEXT_LIKE else kind


def _cell_text(value: object) -> str:
    """The text of a header cell: a date without the zero time, an integral float
    without the ".0", the inner line breaks collapsed into single spaces.

    Without this normalisation a year read from Excel arrives as "2026.0" and a
    two-line caption keeps its "\n" — neither survives being glued into a name.
    """
    if _is_blank(value):
        return ""
    if isinstance(value, dt.datetime):
        if value.hour == value.minute == value.second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def is_header_caption(value: object) -> bool:
    """Whether a header cell holds a caption and not a figure.

    A caption is a word, a date or a year. A figure in the header zone is not a
    title: reports keep check totals above the table, and a header formula whose
    source cell is empty quietly returns 0.

    The type is read from the normalised text rather than from the raw value: a
    year arrives from Excel as 2026.0 and would otherwise pass for a plain integer.
    """
    text = _cell_text(value)
    return bool(text) and detect_type(text) in _CAPTION_TYPES


def detect_typed_header_end(
    ws: Worksheet,
    col: int,
    start_row: int,
    end_row: int,
    merged: dict[tuple[int, int], tuple[int, int]],
) -> int | None:
    """The last header row, determined by the data type change in a column.

    The leading run of cells of one type class is the header; the first cell of
    another class starts the data. Empty rows inside the run do not break it.
    Returns None if the type does not change until the end of the column.
    """
    first_class: str | None = None
    last_header_row: int | None = None

    for row in range(start_row, end_row + 1):
        value = cell_value(ws, row, col, merged)
        if _is_blank(value):
            continue
        current = _type_class(value)
        if first_class is None:
            first_class = current
            last_header_row = row
        elif current == first_class:
            last_header_row = row
        else:
            return last_header_row
    return None


def _row_is_header_row(
    ws: Worksheet,
    row: int,
    extents: dict[int, tuple[int, int]],
    label_col: int,
    merged: dict[tuple[int, int], tuple[int, int]],
) -> bool:
    """Whether the row still belongs to the header, judged outside the label column.

    A header row carries captions and nothing but captions. The data starts at
    the first row that carries a figure — or that carries nothing at all outside
    the label column, which is what a block of row titles without figures looks
    like.

    The label column is skipped on purpose: it holds the row titles, and those
    are text in the header and in the data alike.
    """
    seen_caption = False
    for col in extents:
        if col == label_col:
            continue
        value = cell_value(ws, row, col, merged)
        if _is_blank(value):
            continue
        if not is_header_caption(value):
            return False          # a figure: the data has started
        seen_caption = True
    return seen_caption


def detect_header_layout(ws: Worksheet) -> HeaderLayout:
    """Detects the header and the first data row from the structure of the sheet."""
    merged = _merge_lookup(ws)
    extents = column_extents(ws, merged)
    if not extents:
        return HeaderLayout(orientation="unknown")

    non_empty_cols = sorted(extents)

    label_col = min(extents)
    label_first, label_last = extents[label_col]
    top_row = min(first for first, _ in extents.values())
    bottom_row = max(last for _, last in extents.values())

    # Orientation 1: some column starts above the label column.
    if top_row < label_first:
        # The header may reach one row lower than the label column starts: the
        # bottom header row is often shared by all columns ("Товар | План | Факт")
        # and the label column has its own cell exactly there. Such a row is
        # still a header — but only while it holds captions. A row without
        # captions outside the label column ends the header even when there are
        # no figures in it yet.
        data_start = label_first
        while data_start <= bottom_row and _row_is_header_row(
            ws, data_start, extents, label_col, merged
        ):
            data_start += 1
        if data_start > bottom_row:
            data_start = label_first   # captions all the way down — do not stretch

        header_rows = list(range(top_row, data_start))
        # If the label column gained its own cell in the extended header, the
        # sheet name is no longer needed — it has a title of its own.
        label_has_header = any(
            not _is_blank(cell_value(ws, row, label_col, merged)) for row in header_rows
        )
        return HeaderLayout(
            orientation="staggered",
            header_rows=header_rows,
            data_start_row=data_start,
            data_end_row=bottom_row,
            label_col=label_col,
            label_has_header=label_has_header,
            label_name_from_sheet=not label_has_header,
            non_empty_cols=non_empty_cols,
        )

    # Orientation 2: all columns start together — the type change sets the boundary.
    header_end = detect_typed_header_end(ws, label_col, label_first, label_last, merged)
    if header_end is None or header_end >= bottom_row:
        return HeaderLayout(
            orientation="unknown",
            data_end_row=bottom_row,
            label_col=label_col,
            # The header rows are undetermined, but the label column starts at
            # the top row together with the rest — whatever the header is, it
            # covers the label column too.
            label_has_header=label_first <= top_row,
            label_name_from_sheet=label_first > top_row,
            non_empty_cols=non_empty_cols,
        )

    return HeaderLayout(
        orientation="typed",
        header_rows=list(range(label_first, header_end + 1)),
        data_start_row=header_end + 1,
        data_end_row=bottom_row,
        label_col=label_col,
        label_has_header=True,            # the header is read from the label column itself
        label_name_from_sheet=False,      # the name comes from the header cells
        non_empty_cols=non_empty_cols,
    )



def header_stacks(
    ws: Worksheet, layout: HeaderLayout, max_cols: int = MAX_HEADER_COLS
) -> dict[int, list[str]]:
    """{column: [the values of its header cells, top to bottom]}.

    Merged cells are expanded: the top level of a header is stored only in the
    anchor of the merge range, and without expansion the columns to its right
    would be left without it.

    An upper level written without a merge range is spanned by hand: a value
    extends to the right over every following column and stops only at the next
    column that has a value of its own in that row. How deep the header of a
    spanned column is does not matter — a column of one or two cells, or with no
    header cells at all, belongs to the group just as much. A gap in the column
    numbering breaks the spanning: an entirely empty column separates blocks of
    a report.

    Empty cells are skipped and consecutive repeats are collapsed. A cell that
    holds a figure is skipped as well: it is not a title, and a column whose
    header row happens to carry one behaves as if the cell were empty.
    """
    if not layout.header_rows:
        return {}

    merged = _merge_lookup(ws)
    cols = layout.non_empty_cols[:max_cols]

    stacks: dict[int, list[str]] = {col: [] for col in cols}
    for row in layout.header_rows:
        spanned: str | None = None       # the value of the group open to the left
        previous_col: int | None = None
        for col in cols:
            if previous_col is not None and col != previous_col + 1:
                spanned = None           # an empty column: a new block
            previous_col = col

            value = cell_value(ws, row, col, merged)
            text = _cell_text(value) if is_header_caption(value) else ""
            if text:
                spanned = text           # its own caption: a new group starts here
            else:
                text = spanned or ""     # inside the group open to the left

            parts = stacks[col]
            if text and (not parts or parts[-1] != text):
                parts.append(text)
    return stacks


def glue_header_parts(parts: list[str]) -> str:
    """The header cells of one column glued into a chain, top to bottom.

    Gluing is the job of the code, not of the model: the model is given a ready
    chain and only shortens it into a name. The join is mechanical — the levels
    are made SQL-safe and separated by `__`, nothing is thrown away.
    """
    cleaned = [part for part in (sanitize_name_part(p) for p in parts) if part]
    if not cleaned:
        return ""

    chain = HEADER_LEVEL_SEPARATOR.join(cleaned)
    # The «a name does not start with a digit» rule applies to the whole chain.
    return f"col_{chain}" if chain[0].isdigit() else chain


def glued_headers(
    ws: Worksheet, layout: HeaderLayout, max_cols: int = MAX_HEADER_COLS
) -> dict[int, str]:
    """{column: its header cells glued into one chain}."""
    return {
        col: glue_header_parts(parts)
        for col, parts in header_stacks(ws, layout, max_cols).items()
    }


def format_header_stacks(stacks: dict[int, list[str]]) -> str:
    """A prompt block: the glued chain of every column.

        C: север__2026__q1__план
        A: (ячеек шапки нет)
    """
    lines = []
    for col, parts in stacks.items():
        letter = get_column_letter(col)
        chain = glue_header_parts(parts)
        lines.append(f"{letter}: {chain}" if chain else f"{letter}: (ячеек шапки нет)")
    return "\n".join(lines)
