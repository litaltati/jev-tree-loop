"""Pydantic models of the sheet JSON schema + name sanitization helpers for SQLite."""

from __future__ import annotations

import re

from pydantic import BaseModel, Field, field_validator, model_validator

# Allowed column types. In SQLite, DATE is physically stored as TEXT (ISO-8601),
# see extractor._SQL_TYPES.
ALLOWED_DTYPES = ("TEXT", "REAL", "INTEGER", "DATE")

# Synonyms the LLM may return instead of the canonical types.
_DTYPE_ALIASES = {
    "STRING": "TEXT",
    "STR": "TEXT",
    "VARCHAR": "TEXT",
    "CHAR": "TEXT",
    "NUMERIC": "REAL",
    "NUMBER": "REAL",
    "FLOAT": "REAL",
    "DOUBLE": "REAL",
    "DECIMAL": "REAL",
    "CURRENCY": "REAL",
    "PERCENTAGE": "REAL",
    "PERCENT": "REAL",
    "INT": "INTEGER",
    "BIGINT": "INTEGER",
    "YEAR": "INTEGER",
    "BOOL": "INTEGER",
    "BOOLEAN": "INTEGER",
    "DATETIME": "DATE",
    "TIMESTAMP": "DATE",
    "TIME": "TEXT",
}

_SQL_KEYWORDS = {
    "select", "from", "where", "table", "index", "order", "group", "by",
    "insert", "update", "delete", "create", "drop", "join", "values",
}


def sanitize_identifier(name: str, *, fallback: str = "col") -> str:
    """Converts an arbitrary string into a safe SQLite identifier.

    Spaces -> `_`, special characters are removed, everything is lowercased.
    Cyrillic is preserved: SQLite handles UTF-8 identifiers correctly, and
    transliteration would only hurt the readability of Russian reports.
    """
    text = (name or "").strip().lower()
    text = re.sub(r"\s+", "_", text)
    # Keep letters (of any alphabet), digits and the underscore.
    text = re.sub(r"[^\w]", "_", text, flags=re.UNICODE)
    text = re.sub(r"_+", "_", text).strip("_")
    if not text:
        text = fallback
    if text[0].isdigit():
        text = f"{fallback}_{text}"
    if text in _SQL_KEYWORDS:
        text = f"{text}_"
    return text


def sanitize_table_name(sheet_name: str) -> str:
    """SQLite table name = sanitized sheet name (1 sheet = 1 table)."""
    return sanitize_identifier(sheet_name, fallback="sheet")


def sanitize_name_part(name: str) -> str:
    """Sanitizes a single level of a composite name.

    Unlike sanitize_column_name, it does not add a prefix to a name that
    starts with a digit: a level such as "2026" is a legitimate part inside
    `север__2026__q1`, while the first-character rule applies to the whole name.
    """
    text = (name or "").strip().lower()
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"[^\w]", "_", text, flags=re.UNICODE)
    return re.sub(r"_+", "_", text).strip("_")


def sanitize_column_name(name: str, *, index: int = 0) -> str:
    """SQLite column name. index is used as a fallback for empty headers."""
    return sanitize_identifier(name, fallback=f"col_{index}" if index else "col")


def dedupe_column_names(names: list[str]) -> list[str]:
    """Guarantees uniqueness of column names (`sum`, `sum_2`, `sum_3`, ...)."""
    seen: dict[str, int] = {}
    result: list[str] = []
    for name in names:
        if name not in seen:
            seen[name] = 1
            result.append(name)
        else:
            seen[name] += 1
            result.append(f"{name}_{seen[name]}")
    return result


class ColumnSchema(BaseModel):
    original_col_letter: str       # "A", "B", "C"...
    name: str                      # Column name for SQLite (taken from the header)
    dtype: str                     # "TEXT" | "REAL" | "INTEGER" | "DATE"
    is_row_label: bool = False     # True if this column holds row labels (hierarchy)
    hierarchy_level: int | None = None  # 0 = top level, 1 = sublevel, ...

    @field_validator("original_col_letter", mode="before")
    @classmethod
    def _normalize_letter(cls, v: object) -> str:
        """Accepts "A", "a", "$A$1", as well as a numeric column index."""
        if isinstance(v, int):
            from openpyxl.utils import get_column_letter

            return get_column_letter(v)
        text = str(v or "").strip().upper()
        match = re.search(r"[A-Z]{1,3}", text)
        if not match:
            raise ValueError(f"Не удалось распознать букву колонки: {v!r}")
        return match.group(0)

    @field_validator("dtype", mode="before")
    @classmethod
    def _normalize_dtype(cls, v: object) -> str:
        text = str(v or "TEXT").strip().upper()
        text = _DTYPE_ALIASES.get(text, text)
        return text if text in ALLOWED_DTYPES else "TEXT"

    @field_validator("name", mode="before")
    @classmethod
    def _normalize_name(cls, v: object) -> str:
        # An empty name is not substituted here: SheetSchema will supply a
        # fallback with the column number (col_3) so such columns do not merge.
        return str(v or "").strip()


class SheetSchema(BaseModel):
    sheet_name: str
    table_name: str                # sanitized name for SQLite
    header_rows: list[int] = Field(default_factory=list)  # original Excel coordinates
    data_start_row: int            # first data row (original)
    data_end_row: int              # last data row (original)
    columns: list[ColumnSchema]
    notes: str = ""                # any observations of the model about the structure

    # Not part of the minimal specification, but the prompt asks the model for
    # the column range — we keep it so nothing is lost (not needed in extractor:
    # there the source of truth is the columns list).
    data_start_col: str | None = None
    data_end_col: str | None = None

    @field_validator("data_start_col", "data_end_col", mode="before")
    @classmethod
    def _normalize_col(cls, v: object) -> str | None:
        if v is None or v == "":
            return None
        if isinstance(v, int):
            from openpyxl.utils import get_column_letter

            return get_column_letter(v)
        match = re.search(r"[A-Z]{1,3}", str(v).strip().upper())
        return match.group(0) if match else None

    @model_validator(mode="after")
    def _validate_ranges(self) -> "SheetSchema":
        if not self.columns:
            raise ValueError("Схема должна содержать хотя бы одну колонку")
        if self.data_start_row < 1:
            raise ValueError("data_start_row должен быть >= 1")
        if self.data_end_row < self.data_start_row:
            raise ValueError("data_end_row должен быть >= data_start_row")

        # Column names are made SQL-safe and unique.
        names = [
            sanitize_column_name(c.name, index=i + 1)
            for i, c in enumerate(self.columns)
        ]
        for column, name in zip(self.columns, dedupe_column_names(names)):
            column.name = name

        if not self.table_name:
            self.table_name = sanitize_table_name(self.sheet_name)
        return self
