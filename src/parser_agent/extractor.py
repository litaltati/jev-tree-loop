"""Stage 3: deterministic reading of Excel by the schema and writing to SQLite (no LLM)."""

from __future__ import annotations

import datetime as dt
import logging
import re
import sqlite3
from dataclasses import dataclass, replace
from typing import Any, Callable, Iterable

from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

from .compressor.aggregator import detect_type
from .compressor.converter import prepared_workbook
from .layout import (
    HeaderLayout,
    detect_header_layout,
    glue_header_parts,
    glued_headers,
    is_header_caption,
)
from .models import SheetSchema, dedupe_column_names, sanitize_column_name

logger = logging.getLogger(__name__)

# SQLite has no separate DATE type. Such columns are declared TEXT and stored as
# ISO-8601: a declared DATE type gives NUMERIC affinity, because of which "2024"
# would be silently turned into a number.
_SQL_TYPES = {"TEXT": "TEXT", "REAL": "REAL", "INTEGER": "INTEGER", "DATE": "TEXT"}

# A safety limit against a merge spanning the whole sheet.
_MAX_MERGED_CELLS = 200_000

# Characters that may be stripped before parsing a number: currency, percent,
# a non-breaking space and a space used as a thousands separator.
_CURRENCY_CHARS = "$€₽£¥"
_NUMERIC_RE = re.compile(r"[+-]?\d+(\.\d+)?([eE][+-]?\d+)?$")


def quote_identifier(name: str) -> str:
    """Escapes an identifier for SQL."""
    return '"' + str(name).replace('"', '""') + '"'


def build_merged_lookup(ws: Worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    """`(row, col) -> (anchor_row, anchor_col)` for every merged range.

    Needed so that merged cells (hierarchical row labels) read the value of the
    top-left cell of the range instead of None.
    """
    lookup: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        height = rng.max_row - rng.min_row + 1
        width = rng.max_col - rng.min_col + 1
        if height * width > _MAX_MERGED_CELLS:
            logger.warning("Пропущен слишком большой merged-диапазон %s", rng)
            continue
        anchor = (rng.min_row, rng.min_col)
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                lookup[(row, col)] = anchor
    return lookup


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _to_number(value: str) -> float | None:
    """"1 234,50 ₽" -> 1234.50. None if this is not a number.

    Only currency, percent and thousands-separator spaces are stripped. Letters
    are NOT cut out: otherwise "Q1" would turn into 1, while a text label that
    ended up in a numeric column must stay text.
    """
    text = value.replace("\xa0", " ").strip()
    if not text:
        return None

    text = text.rstrip("%").strip()
    text = text.strip(_CURRENCY_CHARS).strip()
    for token in ("руб.", "руб", "USD", "EUR", "RUB"):
        if text.lower().endswith(token.lower()):
            text = text[: -len(token)].strip()
            break
    text = text.replace(" ", "")  # thousands separator

    # European format: the comma is the decimal separator, the dot the thousands.
    if "," in text and "." in text:
        text = (
            text.replace(".", "").replace(",", ".")
            if text.rfind(",") > text.rfind(".")
            else text.replace(",", "")
        )
    elif "," in text:
        text = text.replace(",", ".")

    if not _NUMERIC_RE.fullmatch(text):
        return None
    return float(text)


def coerce_value(value: object, dtype: str) -> Any:
    """Coerces a cell value to the column type. What cannot be converted stays a string."""
    if _is_blank(value):
        return None

    if dtype == "DATE":
        if isinstance(value, dt.datetime):
            if value.hour == value.minute == value.second == 0:
                return value.date().isoformat()
            return value.isoformat(sep=" ")
        if isinstance(value, (dt.date, dt.time)):
            return value.isoformat()
        return str(value).strip()

    if dtype in ("REAL", "INTEGER"):
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, (int, float)):
            number = float(value)
        elif isinstance(value, str):
            parsed = _to_number(value)
            if parsed is None:
                return value.strip()  # junk in a numeric column is kept as text
            number = parsed
            if "%" in value:
                number /= 100.0
        elif isinstance(value, (dt.datetime, dt.date, dt.time)):
            return value.isoformat()
        else:
            return str(value)
        return int(round(number)) if dtype == "INTEGER" else number

    # TEXT
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.hour == value.minute == value.second == 0 else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() if isinstance(value, str) else str(value)


def apply_layout(schema: SheetSchema, layout: HeaderLayout) -> SheetSchema:
    """Replaces the header and data boundaries with those detected from the structure.

    This is exactly where the model errs: it takes not all header rows, because
    of which part of the header leaks into the data. Structure parsing is
    deterministic, so it takes priority. The end of the data is left as the model
    set it, in case it narrowed the range deliberately (dropping a totals block,
    for example).
    """
    if not layout.detected:
        return schema

    data_end = schema.data_end_row
    if layout.data_end_row is not None:
        data_end = min(data_end, layout.data_end_row)
    data_end = max(data_end, layout.data_start_row)

    return schema.model_copy(
        update={
            "header_rows": list(layout.header_rows),
            "data_start_row": layout.data_start_row,
            "data_end_row": data_end,
        }
    )


def drop_empty_columns(schema: SheetSchema, layout: HeaderLayout) -> SheetSchema:
    """Removes from the schema the columns that are completely empty in Excel.

    Empty columns are used in reports to separate blocks; the model often
    includes them in the schema, and columns without a single value appear in
    the database. A column counts as empty when it has no non-empty cells
    anywhere on the sheet — a column with a header but no data is kept.
    """
    if not layout.non_empty_cols:
        return schema

    allowed = set(layout.non_empty_cols)
    kept = []
    dropped = []
    for column in schema.columns:
        # Comparison by index, not by object: identical column descriptions in
        # the schema are equal to each other and would confuse the selection.
        if column_index_from_string(column.original_col_letter) in allowed:
            kept.append(column)
        else:
            dropped.append(column.original_col_letter)

    if not kept:
        # Nothing is left — the parsing clearly diverged from the schema, do not risk it.
        logger.warning(
            "[%s] все колонки схемы оказались пустыми, фильтр не применён",
            schema.sheet_name,
        )
        return schema

    if dropped:
        logger.info(
            "[%s] пропущены пустые колонки Excel: %s", schema.sheet_name, ", ".join(dropped)
        )
    return schema.model_copy(update={"columns": kept})


def label_column_indexes(schema: SheetSchema, layout: HeaderLayout | None) -> set[int]:
    """Indexes of the row label columns: from the layout parsing and the model flags.

    There can be several labels: a hierarchy is often laid out across adjacent
    columns («Раздел | Статья | Сумма»).
    """
    indexes = {
        column_index_from_string(column.original_col_letter)
        for column in schema.columns
        if column.is_row_label
    }
    if layout is not None and layout.label_col is not None:
        indexes.add(layout.label_col)
    return indexes


def label_column_has_header(
    ws: Worksheet, schema: SheetSchema, layout: HeaderLayout | None = None
) -> bool:
    """Whether the row label column has header cells of its own.

    This is what tells a tree of report items from an ordinary text column.
    Nothing above the label column — the levels to the left are its hierarchy.
    A header above it — it is a plain data column, there is nothing to expand.

    In the typed orientation the answer is always yes: the header boundary is
    read from the label column itself, so its own cells are part of the header —
    such a sheet gets no level columns at all.

    Without a recognized orientation the feature is read straight from the cells,
    so the check works on any sheet.
    """
    if layout is not None and layout.detected:
        return layout.label_has_header

    label_col = layout.label_col if layout is not None else None
    if label_col is None:
        indexes = label_column_indexes(schema, None)
        label_col = min(indexes) if indexes else None
    if label_col is None:
        return False

    merged = build_merged_lookup(ws)
    return any(
        not _is_blank(ws.cell(*merged.get((row, label_col), (row, label_col))).value)
        for row in schema.header_rows
    )


def compose_column_name(
    ws: Worksheet,
    col_idx: int,
    header_rows: list[int],
    merged: dict[tuple[int, int], tuple[int, int]],
) -> str:
    """A column name glued together from the values of all header rows.

    Merged cells are expanded through the anchor of the merge range: the value of
    the top level is stored only in its top-left cell, and without expansion that
    level would be lost. Consecutive identical levels are collapsed. A cell that
    holds a figure instead of a caption is skipped, by the same rule the header
    parsing uses — the two paths must not diverge.
    """
    parts: list[str] = []
    for row in sorted(header_rows):
        cell_row, cell_col = merged.get((row, col_idx), (row, col_idx))
        value = ws.cell(row=cell_row, column=cell_col).value
        if not is_header_caption(value):
            continue
        text = str(coerce_value(value, "TEXT")).strip()
        if text and (not parts or parts[-1] != text):
            parts.append(text)

    # The same gluing as the one the model is shown, so that the chain in the
    # prompt and the chain in the database cannot diverge.
    return glue_header_parts(parts)


def resolve_column_names(
    ws: Worksheet,
    schema: SheetSchema,
    compose_names: bool = False,
    layout: HeaderLayout | None = None,
) -> list[str]:
    """The final column names of the schema.

    By default the names come from the model: it sees the content of the header
    cells and joins the levels meaningfully, which a mechanical concatenation
    cannot do. The code only makes the name SQL-safe (that is done by
    SheetSchema).

    There is one exception, set by the sheet layout: if there is not a single
    header cell above the label column, there is nothing to join — it is named
    after the sheet.

    compose_names=True returns the mechanical concatenation of the header cells
    with `__`: a fallback path for when the model names are unsatisfactory.
    """
    names = [column.name for column in schema.columns]

    merged = build_merged_lookup(ws)
    label_col = layout.label_col if layout is not None else None
    name_from_sheet = bool(layout is not None and layout.label_name_from_sheet)

    # Without a parsed layout, the label column is the one the model marked, and
    # the sheet name is used only if there are no header cells above it.
    if label_col is None:
        for column in schema.columns:
            if column.is_row_label:
                label_col = column_index_from_string(column.original_col_letter)
                name_from_sheet = not any(
                    not _is_blank(
                        ws.cell(
                            *merged.get((row, label_col), (row, label_col))
                        ).value
                    )
                    for row in schema.header_rows
                )
                break

    for index, column in enumerate(schema.columns):
        col_idx = column_index_from_string(column.original_col_letter)
        if col_idx == label_col and name_from_sheet:
            names[index] = sanitize_column_name(schema.sheet_name)
            continue
        if not compose_names or not schema.header_rows:
            continue
        composed = compose_column_name(ws, col_idx, schema.header_rows, merged)
        if composed:
            names[index] = composed
    return names


@dataclass
class HierarchyPlan:
    """The expansion plan of one row label column into level columns."""

    schema_index: int          # the position of the column in schema.columns
    base_name: str             # the name of the label column (the label itself)
    depth: int                 # how many nesting levels were found
    levels: dict[int, int]     # Excel row number -> level (0 = top)

    @property
    def ancestor_names(self) -> list[str]:
        """The names of the ancestor columns: `статья_l0`, `статья_l1`, ..."""
        return [f"{self.base_name}_l{i}" for i in range(self.depth - 1)]


def detect_hierarchy_levels(
    ws: Worksheet,
    col_idx: int,
    start_row: int,
    end_row: int,
    merged: dict[tuple[int, int], tuple[int, int]] | None = None,
) -> dict[int, int]:
    """The nesting level of every row, from two machine-readable Excel features.

    1. `outlineLevel` — row grouping («+»/«-» to the left of the numbers). An
       explicit markup, if the author of the file used it.
    2. `alignment.indent` — the text indent inside a cell. Hierarchies are drawn
       this way more often.

    The `(outlineLevel, indent)` pairs are sorted and numbered: 0, 1, 2... The
    features work together, so a file may use either of them or both at once.

    Only the features to which Excel gives structural meaning are used. Typeface
    (bold, italic) is not used: it is styling without fixed semantics — sections,
    totals and plain emphasis are all marked with it.
    """
    merged = merged or {}
    keys: dict[int, tuple[int, int]] = {}
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        value = cell.value
        source_row = row
        if _is_blank(value) and (row, col_idx) in merged:
            anchor_row, anchor_col = merged[(row, col_idx)]
            cell = ws.cell(row=anchor_row, column=anchor_col)
            value = cell.value
            source_row = anchor_row
        if _is_blank(value):
            continue  # for an empty cell the features mean nothing
        outline = int(ws.row_dimensions[source_row].outlineLevel or 0)
        keys[row] = (outline, int(cell.alignment.indent or 0))

    ranks = {key: level for level, key in enumerate(sorted(set(keys.values())))}
    return {row: ranks[key] for row, key in keys.items()}


def plan_hierarchy(
    ws: Worksheet,
    schema: SheetSchema,
    enabled: bool = True,
    names: list[str] | None = None,
    layout: HeaderLayout | None = None,
) -> list[HierarchyPlan]:
    """Expansion plans for every column with is_row_label=True.

    Levels are expanded only on the sheets where nesting happens at all: where
    the label column has no header of its own (label_column_has_header). With a
    header of its own it is a plain data column with a flat list of values, and
    level columns would only breed emptiness.

    A column is expanded only if there really is more than one level — otherwise
    meaningless columns duplicating the label would appear.
    """
    if not enabled or label_column_has_header(ws, schema, layout):
        return []

    merged = build_merged_lookup(ws)
    max_row = min(schema.data_end_row, ws.max_row or schema.data_end_row)
    names = names or [column.name for column in schema.columns]

    plans: list[HierarchyPlan] = []
    for index, column in enumerate(schema.columns):
        if not column.is_row_label:
            continue
        col_idx = column_index_from_string(column.original_col_letter)
        levels = detect_hierarchy_levels(ws, col_idx, schema.data_start_row, max_row, merged)
        depth = (max(levels.values()) + 1) if levels else 0
        if depth > 1:
            plans.append(
                HierarchyPlan(
                    schema_index=index, base_name=names[index], depth=depth, levels=levels
                )
            )
    return plans


def header_chain_map(
    ws: Worksheet,
    schema: SheetSchema,
    plans: list[HierarchyPlan],
    names: list[str],
    layout: HeaderLayout,
) -> dict[str, str]:
    """{the final column name: the header chain it was made from}.

    The catalog stage shows the chains to the model: in the name the meaning is
    already squeezed to fit SQL, while the chain still holds the words of the
    sheet. Columns without header cells — the level columns among them — are
    simply absent from the map.
    """
    if not layout.header_rows and schema.header_rows:
        # The orientation was not recognized — the header rows are the ones the
        # model gave. Without this the chains would be lost on every sheet whose
        # layout the structure parsing could not name.
        layout = replace(
            layout,
            header_rows=sorted(schema.header_rows),
            non_empty_cols=layout.non_empty_cols
            or sorted(
                column_index_from_string(column.original_col_letter)
                for column in schema.columns
            ),
        )

    chains = glued_headers(ws, layout)
    if not chains:
        return {}

    final = [name for name, _ in effective_columns(schema, plans, names)]
    by_index = {plan.schema_index: plan for plan in plans}

    mapping: dict[str, str] = {}
    position = 0
    for index, column in enumerate(schema.columns):
        plan = by_index.get(index)
        if plan is not None:
            position += len(plan.ancestor_names)   # the ancestors go before the label
        if position >= len(final):
            break
        chain = chains.get(column_index_from_string(column.original_col_letter), "")
        if chain:
            mapping[final[position]] = chain
        position += 1
    return mapping


def effective_columns(
    schema: SheetSchema,
    plans: list[HierarchyPlan],
    names: list[str] | None = None,
) -> list[tuple[str, str]]:
    """The final list of table columns `(name, SQL type)`, accounting for expansion.

    The order for a label column: ancestors -> the label itself. The level number
    is not written into the table: it is internal, needed only to carry the
    ancestors forward, and in SQL its role is fully served by the ancestor
    columns themselves.
    """
    by_index = {plan.schema_index: plan for plan in plans}
    resolved = names or [column.name for column in schema.columns]
    out_names: list[str] = []
    types: list[str] = []

    for index, column in enumerate(schema.columns):
        plan = by_index.get(index)
        if plan is not None:
            for ancestor in plan.ancestor_names:
                out_names.append(ancestor)
                types.append("TEXT")
        out_names.append(resolved[index])
        types.append(_SQL_TYPES.get(column.dtype, "TEXT"))
    # Expansion and name composition may have created a name clash between columns.
    return list(zip(dedupe_column_names(out_names), types))


def create_table_sql(
    schema: SheetSchema,
    plans: list[HierarchyPlan] | None = None,
    names: list[str] | None = None,
) -> str:
    """`CREATE TABLE IF NOT EXISTS ...` over the schema columns, accounting for expansion."""
    columns = ", ".join(
        f"{quote_identifier(name)} {sql_type}"
        for name, sql_type in effective_columns(schema, plans or [], names)
    )
    return f"CREATE TABLE IF NOT EXISTS {quote_identifier(schema.table_name)} ({columns})"


def _insert_sql(
    schema: SheetSchema,
    plans: list[HierarchyPlan] | None = None,
    names: list[str] | None = None,
) -> str:
    columns = effective_columns(schema, plans or [], names)
    names = ", ".join(quote_identifier(name) for name, _ in columns)
    placeholders = ", ".join("?" for _ in columns)
    return f"INSERT INTO {quote_identifier(schema.table_name)} ({names}) VALUES ({placeholders})"


def has_data_value(values: Iterable[object]) -> bool:
    """Whether at least one of the values is non-text (a number, a date, a percentage)."""
    return any(detect_type(value) not in ("Text", "Email", "Empty") for value in values)


def iter_rows_for_schema(
    ws: Worksheet,
    schema: SheetSchema,
    plans: list[HierarchyPlan] | None = None,
    skip_text_only_rows: bool = True,
    label_indexes: set[int] | None = None,
) -> Iterable[list[Any]]:
    """The data rows of the sheet, coerced to the schema types.

    Rows without values of their own are skipped and, with skip_text_only_rows,
    so are rows for which BOTH conditions hold: the whole row has not a single
    non-text value AND the label column is empty (accounting for merges). A row
    mixing text, numbers and blanks is kept; a row with a label but without
    numbers is kept too — only a dangling row with no anchor counts as empty.

    For hierarchical columns, ancestor columns are added: the label of each level
    is carried down until the next label of the same level is met. The level
    number itself does not reach the result — it only drives the carrying. The
    carry state is updated by skipped rows as well: a section row is exactly what
    sets the ancestor for the rows below it.
    """
    merged = build_merged_lookup(ws)
    col_indexes = [column_index_from_string(col.original_col_letter) for col in schema.columns]
    max_row = min(schema.data_end_row, ws.max_row or schema.data_end_row)

    # The position of the label column: the second skip condition is checked by it.
    labels = label_indexes or set()
    label_positions = [
        position for position, col_idx in enumerate(col_indexes) if col_idx in labels
    ]
    label_position = min(label_positions) if label_positions else 0

    by_index = {plan.schema_index: plan for plan in (plans or [])}
    # The current carry state: plan -> [level 0 label, level 1 label, ...]
    current: dict[int, list[str | None]] = {
        index: [None] * plan.depth for index, plan in by_index.items()
    }

    for row_idx in range(schema.data_start_row, max_row + 1):
        raw_values: list[object] = []
        resolved: list[object] = []
        for col_idx in col_indexes:
            raw = ws.cell(row=row_idx, column=col_idx).value
            raw_values.append(raw)
            if _is_blank(raw) and (row_idx, col_idx) in merged:
                anchor_row, anchor_col = merged[(row_idx, col_idx)]
                raw = ws.cell(row=anchor_row, column=anchor_col).value
            resolved.append(raw)

        # A row counts as empty when it has no values of its own: values
        # "inherited" from a merged header do not bring a row to life.
        blank_row = all(_is_blank(value) for value in raw_values)
        orphan_row = (
            skip_text_only_rows
            and not has_data_value(raw_values)                    # text and blanks only
            # The label is read accounting for merges: a merged cell covers the
            # row, and such a row does not dangle without an anchor.
            and _is_blank(resolved[label_position])
        )

        row: list[Any] = []
        for index, (value, column) in enumerate(zip(resolved, schema.columns)):
            plan = by_index.get(index)
            coerced = coerce_value(value, column.dtype)

            if plan is None:
                row.append(coerced)
                continue

            level = plan.levels.get(row_idx)
            if level is not None:
                current[index][level] = coerced
                # Deeper levels do not belong to the new label.
                for deeper in range(level + 1, plan.depth):
                    current[index][deeper] = None

            row.extend(current[index][: plan.depth - 1])  # ancestors
            row.append(coerced)                            # the label itself

        # The row is skipped AFTER the carry state is updated: a dropped row can
        # still be an ancestor for the rows below it.
        if blank_row or orphan_row:
            continue

        yield row


def extract_to_sqlite(
    wb_path: str,          # path to the source .xlsx
    schema: SheetSchema,   # the schema from the LLM
    db_path: str,          # path to the SQLite file
    *,
    ws: Worksheet | None = None,   # an already open sheet — then the file is not read again
    drop_existing: bool = False,   # recreate the table (idempotent re-run)
    expand_hierarchy: bool = True,  # expand the row label hierarchy into level columns
    compose_names: bool = False,    # compose names in code instead of the model names
    detect_layout: bool = True,     # detect the header boundaries from the sheet structure
    skip_empty_columns: bool = True,     # do not create columns that are empty in Excel
    skip_text_only_rows: bool = True,     # do not store rows without data (text only)
    on_columns: Callable[[dict[str, str]], None] | None = None,  # {column: header chain}
) -> int:                  # the number of inserted rows
    """Reads the sheet by the schema and writes it into a single SQLite table.

    The workbook is opened with data_only=True: a formula arrives as the value
    Excel computed for it, which is what has to end up in the database. The
    caller that already holds the sheet passes it in `ws`, and then the file is
    not read a second time.
    """
    workbook = None
    if ws is None:
        # A non-.xlsx source is converted first (formulas -> computed values).
        with prepared_workbook(wb_path) as source:
            workbook = source.load()
        if schema.sheet_name not in workbook.sheetnames:
            workbook.close()
            raise KeyError(f"Лист {schema.sheet_name!r} не найден в {wb_path}")
        ws = workbook[schema.sheet_name]
    try:

        # The layout is parsed whenever it is allowed: even if the orientation was
        # not recognized, the list of non-empty columns is known from it.
        layout = detect_header_layout(ws) if (detect_layout or skip_empty_columns) else HeaderLayout("unknown")

        if detect_layout and layout.detected:
            logger.info(
                "[%s] раскладка: %s, шапка %s, данные с %d",
                schema.sheet_name, layout.orientation,
                layout.header_rows, layout.data_start_row,
            )
            schema = apply_layout(schema, layout)

        if skip_empty_columns:
            schema = drop_empty_columns(schema, layout)

        names = resolve_column_names(
            ws, schema, compose_names=compose_names,
            layout=layout if (detect_layout and layout.detected) else None,
        )
        for column, name in zip(schema.columns, names):
            if name != column.name:
                logger.debug(
                    "[%s] колонка %s: имя собрано из заголовков -> %s",
                    schema.sheet_name, column.original_col_letter, name,
                )
        plans = plan_hierarchy(
            ws, schema, enabled=expand_hierarchy, names=names, layout=layout
        )
        for plan in plans:
            logger.info(
                "[%s] иерархия в колонке %s: %d уровня(ей) -> +%d колонок предков",
                schema.sheet_name, plan.base_name, plan.depth, plan.depth - 1,
            )
        if on_columns is not None:
            # The chains are known only here, with the sheet still open: the
            # catalog stage writes the column descriptions from them.
            on_columns(header_chain_map(ws, schema, plans, names, layout))

        labels = label_column_indexes(schema, layout if layout.label_col is not None else None)
        rows = list(
            iter_rows_for_schema(
                ws, schema, plans,
                skip_text_only_rows=skip_text_only_rows, label_indexes=labels,
            )
        )
        if skip_text_only_rows and not rows:
            # The table is entirely textual (a directory of names, addresses) —
            # dropping everything in it makes no sense, return the rows as is.
            rows = list(iter_rows_for_schema(ws, schema, plans, skip_text_only_rows=False))
            if rows:
                logger.warning(
                    "[%s] все строки текстовые, фильтр строк не применён", schema.sheet_name
                )
    finally:
        if workbook is not None:
            workbook.close()

    connection = sqlite3.connect(db_path)
    try:
        connection.text_factory = str  # UTF-8 strings
        cursor = connection.cursor()
        if drop_existing:
            cursor.execute(f"DROP TABLE IF EXISTS {quote_identifier(schema.table_name)}")
        cursor.execute(create_table_sql(schema, plans, names))
        if rows:
            cursor.executemany(_insert_sql(schema, plans, names), rows)
        connection.commit()
    finally:
        connection.close()

    logger.info(
        "[%s] вставлено строк: %d -> таблица %s", schema.sheet_name, len(rows), schema.table_name
    )
    return len(rows)
