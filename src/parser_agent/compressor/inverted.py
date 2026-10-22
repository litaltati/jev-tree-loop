"""Module 2: Inverted-Index Translation (SpreadsheetLLM, arXiv:2407.09025).

Instead of listing "address -> value" we build an inverted index
"value -> addresses", collapsing adjacent addresses into ranges. Empty cells
are not written at all — on real sheets this is the main source of savings.

IMPORTANT: addresses are emitted in the ORIGINAL Excel coordinates, not in the
coordinates of the compact matrix. That way the LLM returns the schema directly
in the coordinates of the source file (stage 3 reads the file by them), and no
back-translation is needed. As a consequence, ranges are only collapsed where
the cells are adjacent in the original — a gap caused by dropped rows is not
masked by a range.
"""

from __future__ import annotations

import datetime as dt
import json
from collections import defaultdict
from typing import Iterable

from openpyxl.utils import get_column_letter

from .anchors import CompactSheet

# Long text values are truncated: the beginning of the string is enough to
# recognize the structure, and it saves a noticeable number of tokens.
MAX_VALUE_LEN = 200


def normalize_value(value: object) -> str:
    """Converts a cell value into a compact string (the key of the inverted index)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.hour == value.minute == value.second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:g}"
    text = str(value).strip()
    if len(text) > MAX_VALUE_LEN:
        text = text[:MAX_VALUE_LEN] + "…"
    return text


def _runs(sorted_values: list[int]) -> list[tuple[int, int]]:
    """[1,2,3,7,8] -> [(1,3),(7,8)]"""
    runs: list[tuple[int, int]] = []
    for value in sorted_values:
        if runs and value == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], value)
        else:
            runs.append((value, value))
    return runs


def _rect_address(r1: int, c1: int, r2: int, c2: int) -> str:
    start = f"{get_column_letter(c1)}{r1}"
    if (r1, c1) == (r2, c2):
        return start
    return f"{start}:{get_column_letter(c2)}{r2}"


def fold_addresses(coords: Iterable[tuple[int, int]]) -> str:
    """Collapses a list of (row, col) into a compact range notation.

    First vertical runs inside a column (A1,A2,A3 -> A1:A3), then identical
    runs in adjacent columns are merged into a rectangle
    (A1:A3 + B1:B3 -> A1:B3). The result is comma-separated addresses.
    """
    by_col: dict[int, set[int]] = defaultdict(set)
    for row, col in coords:
        by_col[col].add(row)

    # (row_start, row_end) -> the list of columns with the same run
    cols_by_run: dict[tuple[int, int], list[int]] = defaultdict(list)
    for col, rows in by_col.items():
        for run in _runs(sorted(rows)):
            cols_by_run[run].append(col)

    rects: list[tuple[int, int, int, int]] = []
    for (r1, r2), cols in cols_by_run.items():
        for c1, c2 in _runs(sorted(cols)):
            rects.append((r1, c1, r2, c2))

    rects.sort(key=lambda rect: (rect[1], rect[0]))  # by column, then by row
    return ",".join(_rect_address(*rect) for rect in rects)


def build_inverted_index(compact: CompactSheet) -> dict[str, str]:
    """`{value: "A1:A3,C5"}` over all non-empty cells of the compact matrix."""
    coords_by_value: dict[str, list[tuple[int, int]]] = defaultdict(list)
    for _, _, info in compact.iter_cells():
        text = normalize_value(info.value)
        if not text:
            continue
        coords_by_value[text].append((info.row, info.col))
    return {value: fold_addresses(coords) for value, coords in coords_by_value.items()}


def to_inverted_json(compact: CompactSheet) -> str:
    """The inverted index as a JSON string (output of module 2)."""
    return json.dumps(build_inverted_index(compact), ensure_ascii=False, separators=(",", ":"))
