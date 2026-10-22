"""Combines modules 1 -> 2 -> 3 into a single compress_sheet()."""

from __future__ import annotations

import logging
from dataclasses import dataclass

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .aggregator import aggregate
from .anchors import DEFAULT_K, DEFAULT_THRESHOLD, CompactSheet, build_compact_sheet
from .inverted import normalize_value, to_inverted_json

logger = logging.getLogger(__name__)

# How many merged ranges to show in the header (they matter for the hierarchy
# of row labels, but there can be very many of them).
MAX_MERGED_IN_HEADER = 40

# A sheet with fewer than two non-empty rows contains neither a header nor data.
MIN_NON_EMPTY_ROWS = 2


def estimate_tokens(text: str) -> int:
    """A rough estimate of the number of tokens (1 token ≈ 4 characters)."""
    return len(text) // 4


@dataclass
class CompressionStats:
    sheet_name: str
    original_chars: int
    original_tokens: int
    compressed_chars: int
    compressed_tokens: int
    rows_before: int
    rows_after: int
    cols_before: int
    cols_after: int
    aggregated: bool

    @property
    def ratio(self) -> float:
        """Compression ratio: how many times fewer tokens there are now."""
        return self.original_tokens / self.compressed_tokens if self.compressed_tokens else 0.0

    def summary(self) -> str:
        return (
            f"[{self.sheet_name}] {self.rows_before}x{self.cols_before} -> "
            f"{self.rows_after}x{self.cols_after} | tokens {self.original_tokens} -> "
            f"{self.compressed_tokens} | ratio {self.ratio:.2f}x"
            f"{'' if self.aggregated else ' (без агрегации)'}"
        )


def vanilla_encoding(ws: Worksheet) -> str:
    """The baseline ("vanilla") sheet encoding used to measure the compression ratio.

    The format from the paper: non-empty cells listed as `address,value`,
    separated by `|`. Empty cells are not counted — otherwise the ratio would
    be inflated artificially.
    """
    parts: list[str] = []
    for row in ws.iter_rows():
        for cell in row:
            text = normalize_value(cell.value)
            if text:
                parts.append(f"{get_column_letter(cell.column)}{cell.row},{text}")
    return "|".join(parts)


def non_empty_row_count(ws: Worksheet) -> int:
    """The number of rows that hold at least one non-empty value."""
    count = 0
    for row in ws.iter_rows(values_only=True):
        if any(v is not None and (not isinstance(v, str) or v.strip()) for v in row):
            count += 1
    return count


def _header(compact: CompactSheet, aggregated: bool) -> str:
    """Short metadata about the sheet, placed before the JSON body."""
    kept_rows = compact.kept_rows
    kept_cols = compact.kept_cols
    merged = compact.merged_ranges[:MAX_MERGED_IN_HEADER]
    lines = [
        f"# sheet: {compact.sheet_name}",
        f"# original_size: {compact.original_n_rows} rows x {compact.original_n_cols} cols",
        f"# kept_rows: {_ranges_str(kept_rows)}",
        f"# kept_cols: {_ranges_str(kept_cols, as_letters=True)}",
    ]
    if merged:
        more = "" if len(compact.merged_ranges) <= MAX_MERGED_IN_HEADER else " ..."
        lines.append(f"# merged_cells: {', '.join(merged)}{more}")
    lines.append(
        "# адреса ниже — ОРИГИНАЛЬНЫЕ координаты Excel; "
        + (
            "числа/даты свёрнуты в записи вида <формат>: <диапазон>"
            if aggregated
            else "значения перечислены как есть"
        )
    )
    return "\n".join(lines)


def _ranges_str(values: list[int], as_letters: bool = False) -> str:
    """[1,2,3,7] -> "1-3,7" (or "A-C,G" for columns)."""
    if not values:
        return "-"
    runs: list[tuple[int, int]] = []
    for value in values:
        if runs and value == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], value)
        else:
            runs.append((value, value))

    def fmt(value: int) -> str:
        return get_column_letter(value) if as_letters else str(value)

    return ",".join(fmt(a) if a == b else f"{fmt(a)}-{fmt(b)}" for a, b in runs)


def compress_sheet_with_stats(
    ws: Worksheet,
    *,
    skip_aggregation: bool = False,
    k: int = DEFAULT_K,
    threshold: float = DEFAULT_THRESHOLD,
) -> tuple[str, CompressionStats]:
    """The same as compress_sheet(), but additionally returns the statistics."""
    compact = build_compact_sheet(ws, k=k, threshold=threshold)

    inverted = to_inverted_json(compact)
    body = inverted if skip_aggregation else aggregate(inverted, compact)
    text = f"{_header(compact, aggregated=not skip_aggregation)}\n{body}"

    baseline = vanilla_encoding(ws)
    stats = CompressionStats(
        sheet_name=ws.title,
        original_chars=len(baseline),
        original_tokens=max(estimate_tokens(baseline), 1),
        compressed_chars=len(text),
        compressed_tokens=max(estimate_tokens(text), 1),
        rows_before=compact.original_n_rows,
        rows_after=compact.n_rows,
        cols_before=compact.original_n_cols,
        cols_after=compact.n_cols,
        aggregated=not skip_aggregation,
    )
    logger.info(stats.summary())
    return text, stats


def compress_sheet(
    ws: Worksheet,
    *,
    skip_aggregation: bool = False,
    k: int = DEFAULT_K,
    threshold: float = DEFAULT_THRESHOLD,
) -> str:
    """Applies modules 1→2→3 in sequence and returns the compressed text.

    skip_aggregation=True disables module 3 — handy for comparing the gain.
    """
    text, _ = compress_sheet_with_stats(
        ws,
        skip_aggregation=skip_aggregation,
        k=k,
        threshold=threshold,
    )
    return text
